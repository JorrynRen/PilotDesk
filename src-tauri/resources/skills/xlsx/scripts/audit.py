"""
公式结构审计

audit_formulas(filename): 检测 recalc 发现不了的结构错误
"""

import json
import re
from collections import OrderedDict

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.cell import range_boundaries

_MAX_SAMPLES = 5
_MAX_GROUPS_PER_TYPE = 20

# 长函数名在前，避免 COUNT 误匹配 COUNTIF
_AGG_A1_PATTERN = re.compile(
    r"(SUMIFS|COUNTIFS|AVERAGEIFS|SUMIF|COUNTIF|AVERAGEIF|"
    r"SUM|AVERAGE|COUNTA|COUNT|MIN|MAX)\s*\(\s*"
    r"(?:(?P<sheet>'[^']+'|[^'!]+)!)?"
    r"(?P<c1>[A-Z]+)(?P<r1>\d+):(?P<c2>[A-Z]+)(?P<r2>\d+)",
    re.IGNORECASE,
)

_STRUCTURED_REF_PATTERN = re.compile(
    r"[A-Za-z_\u4e00-\u9fff][A-Za-z0-9_\u4e00-\u9fff]*\[",
)

# 列聚合结构化引用（排除 [@列] 行引用）；表内写此类公式会循环引用
_STRUCTURED_COL_AGG_PATTERN = re.compile(
    r"(SUMIFS|COUNTIFS|AVERAGEIFS|SUMIF|COUNTIF|AVERAGEIF|"
    r"SUM|AVERAGE|COUNTA|COUNT|MIN|MAX)\s*\(\s*"
    r"(?P<table>[A-Za-z_\u4e00-\u9fff][A-Za-z0-9_\u4e00-\u9fff]*)"
    r"\[(?P<col>(?!@)[^\]]+)\]",
    re.IGNORECASE,
)


def audit_formulas(filename):
    """
    检测结构问题（启发式，可能误报）：
    1. 引用范围缺失：SUM/AVG 等聚合公式漏行/漏列
    2. 公式被覆盖为硬编码值：原应含公式的单元格变成裸数字
    3. 硬编码假设藏在公式里：公式含非整数常量
    4. Table 内仍用固定 A1 聚合：应改用结构化引用
    5. Table 内对同表列做结构化合计：会循环引用，合计应放表外
    """
    wb = load_workbook(filename, data_only=False)
    warnings = []

    for ws in wb.worksheets:
        formulas = {}

        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formulas[(cell.row, cell.column)] = cell.value

        warnings.extend(_check_range_gaps(ws, formulas))
        warnings.extend(_check_hardcoded_assumptions(ws, formulas))
        warnings.extend(_check_inconsistent_formulas(ws, formulas))
        warnings.extend(_check_fixed_range_on_table(ws, formulas))
        warnings.extend(_check_structured_ref_inside_table(ws, formulas))

    warnings.extend(_check_overwritten_formulas(wb))

    return _finalize_warnings(warnings)


def _finalize_warnings(warnings):
    raw_count = len(warnings)
    if raw_count == 0:
        return {
            "status": "success",
            "total_warnings": 0,
            "raw_warning_count": 0,
            "suppressed_count": 0,
            "summary": {},
            "warnings": [],
        }

    summary = {}
    for w in warnings:
        summary[w["type"]] = summary.get(w["type"], 0) + 1

    grouped, suppressed = _aggregate_warnings(warnings)
    return {
        "status": "warnings_found",
        "total_warnings": len(grouped),
        "raw_warning_count": raw_count,
        "suppressed_count": suppressed,
        "summary": summary,
        "warnings": grouped,
    }


def _aggregate_warnings(warnings):
    """按类型+关键特征合并；每类最多保留 _MAX_GROUPS_PER_TYPE 组。"""
    groups = OrderedDict()
    for w in warnings:
        key = _warning_group_key(w)
        if key not in groups:
            groups[key] = {
                "type": w["type"],
                "count": 0,
                "samples": [],
                "message": w.get("message", ""),
            }
            for field in ("value", "table", "formula"):
                if field in w:
                    groups[key][field] = w[field]
        g = groups[key]
        g["count"] += 1
        if len(g["samples"]) < _MAX_SAMPLES:
            sample = {"cell": w.get("cell"), "message": w.get("message")}
            if "formula" in w:
                sample["formula"] = w["formula"]
            if "value" in w:
                sample["value"] = w["value"]
            g["samples"].append(sample)
        if g["count"] > 1:
            g["message"] = (
                f"{w.get('message', '')}（共 {g['count']} 处，详见 samples）"
            )

    by_type = OrderedDict()
    for key, g in groups.items():
        by_type.setdefault(g["type"], []).append(g)

    result = []
    suppressed = 0
    for _type, items in by_type.items():
        if len(items) > _MAX_GROUPS_PER_TYPE:
            kept = items[:_MAX_GROUPS_PER_TYPE]
            for extra in items[_MAX_GROUPS_PER_TYPE:]:
                suppressed += extra["count"]
            result.extend(kept)
        else:
            result.extend(items)
    return result, suppressed


def _warning_group_key(w):
    t = w["type"]
    if t == "hardcoded_assumption":
        return (t, w.get("value"))
    if t == "possible_overwrite":
        cell = w.get("cell") or ""
        sheet, _, ref = cell.partition("!")
        col_m = re.match(r"[A-Z]+", ref or "")
        return (t, sheet, col_m.group(0) if col_m else ref)
    if t == "fixed_range_on_table":
        return (t, w.get("table"), w.get("formula"))
    if t == "structured_ref_inside_table":
        return (t, w.get("table"), w.get("formula"))
    if t == "range_gap":
        return (t, w.get("formula"))
    if t == "inconsistent_formula":
        return (t, re.sub(r"\d+", "N", w.get("formula") or ""))
    return (t, w.get("cell"), w.get("message"))


def _is_numeric_data_cell(value) -> bool:
    """上方单元格是否像被漏掉的数值数据（表头文字不算）。"""
    if value is None or isinstance(value, bool):
        return False
    if isinstance(value, (int, float)):
        return True
    if isinstance(value, str):
        if value.startswith("="):
            return False
        s = value.strip().replace(",", "")
        try:
            float(s)
            return True
        except ValueError:
            return False
    # datetime 等可聚合类型视为数据
    return True


def _check_range_gaps(ws, formulas):
    """检测 SUM/AVG/COUNT 等聚合公式的引用范围是否可能漏行"""
    warnings = []
    agg_pattern = re.compile(
        r"=(SUM|AVERAGE|COUNT|COUNTA|MIN|MAX)\s*\(\s*([A-Z]+)(\d+):([A-Z]+)(\d+)\s*\)",
        re.IGNORECASE,
    )

    for (row, col), formula in formulas.items():
        match = agg_pattern.search(formula)
        if not match:
            continue

        func_name = match.group(1).upper()
        start_col, start_row = match.group(2).upper(), int(match.group(3))
        end_col, end_row = match.group(4).upper(), int(match.group(5))

        # 同行横向聚合（如 SUM(B6:G6)）不做「上方漏行」检查
        if start_row == end_row:
            continue

        if start_row > 2:
            check_row = start_row - 1
            check_col_idx = column_index_from_string(start_col)
            cell_above = ws.cell(row=check_row, column=check_col_idx)
            if _is_numeric_data_cell(cell_above.value):
                cell_ref = f"{start_col}{start_row}:{end_col}{end_row}"
                warnings.append(
                    {
                        "type": "range_gap",
                        "cell": f"{ws.title}!{get_column_letter(col)}{row}",
                        "formula": formula,
                        "message": (
                            f"{func_name} 范围 {cell_ref} 可能漏掉了第 {check_row} 行的数据"
                            f"（{start_col}{check_row}={cell_above.value}）"
                        ),
                    }
                )

    return warnings


def _axis_numeric_values(ws):
    """收集可能作为矩阵轴标签的裸数值（同行/同列检索用）。"""
    by_row = {}
    by_col = {}
    max_row = ws.max_row or 1
    max_col = ws.max_column or 1
    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            val = ws.cell(row=r, column=c).value
            if isinstance(val, (int, float)) and not isinstance(val, bool):
                by_row.setdefault(r, set()).add(float(val))
                by_col.setdefault(c, set()).add(float(val))
    return by_row, by_col


def _value_on_axis(by_row, by_col, row, col, val):
    """公式单元格的同行或同列轴上是否出现同值裸数字。"""
    target = float(val)
    if target in by_row.get(row, ()):
        return True
    if target in by_col.get(col, ()):
        return True
    return False


def _check_hardcoded_assumptions(ws, formulas):
    """检测公式中藏有非整数常量（可能是业务假设）"""
    warnings = []
    decimal_pattern = re.compile(r"(?<![\w.])\d+\.\d+(?![\w])")
    # 仅小数白名单（整数不会被 decimal_pattern 匹配）
    whitelist = {0.5, 0.01, 0.001}
    by_row, by_col = _axis_numeric_values(ws)

    for (row, col), formula in formulas.items():
        matches = decimal_pattern.findall(formula)
        for match in matches:
            val = float(match)
            if val in whitelist:
                continue
            if val == int(val):
                continue
            if _value_on_axis(by_row, by_col, row, col, val):
                continue
            warnings.append(
                {
                    "type": "hardcoded_assumption",
                    "cell": f"{ws.title}!{get_column_letter(col)}{row}",
                    "formula": formula,
                    "value": val,
                    "message": (
                        f"公式含常量 {val}，可能是业务假设（税率/增长率/折现率等），"
                        "建议提取到有标签的独立单元格"
                    ),
                }
            )

    return warnings


def _check_inconsistent_formulas(ws, formulas):
    """检测同一列中结构异常的公式"""
    warnings = []
    col_formulas = {}
    for (row, col), formula in formulas.items():
        if col not in col_formulas:
            col_formulas[col] = []
        col_formulas[col].append((row, formula))

    for col, items in col_formulas.items():
        if len(items) < 3:
            continue

        agg_prefix = re.compile(
            r"^=(SUM|AVERAGE|COUNT|COUNTA|MIN|MAX|SUBTOTAL)\s*\(",
            re.IGNORECASE,
        )
        items = [(row, formula) for row, formula in items if not agg_prefix.match(formula)]
        if len(items) < 3:
            continue

        def normalize(f):
            # 字符串条件不同（如 COUNTIF 分组名）视为同构，避免误报
            f = re.sub(r'"[^"]*"', '"S"', f)
            f = re.sub(r"'[^']*'", "'S'", f)
            return re.sub(r"\d+", "N", f)

        patterns = {}
        for row, formula in items:
            norm = normalize(formula)
            if norm not in patterns:
                patterns[norm] = []
            patterns[norm].append(row)

        if len(patterns) > 1:
            sorted_patterns = sorted(patterns.items(), key=lambda x: len(x[1]))
            minority_pattern, minority_rows = sorted_patterns[0]
            if len(minority_rows) <= 2 and len(patterns) > 1:
                col_letter = get_column_letter(col)
                items_dict = dict(items)
                for r in minority_rows:
                    orig_formula = items_dict.get(r, "")
                    warnings.append(
                        {
                            "type": "inconsistent_formula",
                            "cell": f"{ws.title}!{col_letter}{r}",
                            "formula": orig_formula,
                            "message": "同列公式结构与其他行不一致，可能是引用错误",
                        }
                    )

    return warnings


def _is_axis_label_cell(ws, row, col, col_formulas):
    """
    多列公式块之间的水平数字轴：至少 2 列在该行上下均有公式，
    且这些列在该行多为裸数字 → 视为矩阵/参数轴，而非单列公式被覆盖。
    """
    matrix_cols = [
        c
        for c, frows in col_formulas.items()
        if any(r < row for r in frows) and any(r > row for r in frows)
    ]
    if len(matrix_cols) < 2 or col not in matrix_cols:
        return False
    numeric_in_strip = 0
    for c in matrix_cols:
        val = ws.cell(row=row, column=c).value
        if isinstance(val, (int, float)) and not isinstance(val, bool):
            numeric_in_strip += 1
    return numeric_in_strip >= 2


def _cell_referenced_in_formulas(row, col, formulas):
    """裸数字若被其它公式引用，视为输入参数而非公式被覆盖。"""
    letter = get_column_letter(col)
    pattern = re.compile(
        rf"(?<![A-Z])\$?{re.escape(letter)}\$?{row}(?!\d)",
        re.IGNORECASE,
    )
    for (fr, fc), formula in formulas.items():
        if fr == row and fc == col:
            continue
        if pattern.search(formula):
            return True
    return False


def _is_formula_column_edge_literal(row, formula_rows):
    """
    公式列紧邻块首/块尾的裸数字：上方或下方有连续公式块（≥2 行），
    且该行恰好贴在块边界外一行。覆盖「整列写公式后末行被改成硬编码」类漏检。
    """
    if len(formula_rows) < 2:
        return False
    min_f, max_f = min(formula_rows), max(formula_rows)
    return row == max_f + 1 or row == min_f - 1


def _check_overwritten_formulas(wb):
    """同列公式块中间或紧邻首尾的裸数值，疑似公式被覆盖。"""
    warnings = []
    for ws in wb.worksheets:
        col_formulas = {}
        col_numerics = {}
        formulas = {}
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    col_formulas.setdefault(cell.column, set()).add(cell.row)
                    formulas[(cell.row, cell.column)] = cell.value
                elif isinstance(cell.value, (int, float)) and not isinstance(
                    cell.value, bool
                ):
                    col_numerics.setdefault(cell.column, {})[cell.row] = cell.value

        for col, numerics in col_numerics.items():
            formula_rows = col_formulas.get(col)
            if not formula_rows:
                continue
            for row, val in numerics.items():
                has_above = any(r < row for r in formula_rows)
                has_below = any(r > row for r in formula_rows)
                is_middle_hole = has_above and has_below
                is_edge = _is_formula_column_edge_literal(row, formula_rows)
                if not (is_middle_hole or is_edge):
                    continue
                if _is_axis_label_cell(ws, row, col, col_formulas):
                    continue
                if _cell_referenced_in_formulas(row, col, formulas):
                    continue
                if is_middle_hole:
                    loc = "同列公式之间（上方/下方均有公式）"
                elif row > max(formula_rows):
                    loc = "同列公式块末尾下一行"
                else:
                    loc = "同列公式块开头上一行"
                warnings.append(
                    {
                        "type": "possible_overwrite",
                        "cell": f"{ws.title}!{get_column_letter(col)}{row}",
                        "value": val,
                        "message": (
                            f"数值 {val} 位于{loc}，可能是公式被覆盖为硬编码值"
                        ),
                    }
                )

    return warnings


def _parse_table_bounds(table):
    min_col, min_row, max_col, max_row = range_boundaries(table.ref)
    return min_row, max_row, min_col, max_col


def _ranges_overlap(a, b):
    a_min_r, a_max_r, a_min_c, a_max_c = a
    b_min_r, b_max_r, b_min_c, b_max_c = b
    return not (
        a_max_r < b_min_r
        or b_max_r < a_min_r
        or a_max_c < b_min_c
        or b_max_c < a_min_c
    )


def _normalize_sheet_name(name):
    if not name:
        return None
    name = name.strip()
    if name.startswith("'") and name.endswith("'"):
        name = name[1:-1].replace("''", "'")
    return name


def _iter_table_bounds(ws):
    """Yield (display_name, ref, bounds) for each Table on the sheet."""
    for table in list(getattr(ws, "tables", {}).values()):
        try:
            name = getattr(table, "displayName", None) or table.name
            yield name, table.ref, _parse_table_bounds(table)
        except Exception:
            continue


def _cell_in_bounds(row, col, bounds) -> bool:
    min_row, max_row, min_col, max_col = bounds
    return min_row <= row <= max_row and min_col <= col <= max_col


def _check_fixed_range_on_table(ws, formulas):
    """Table 区域内仍用 A1 固定范围聚合 → 应改用结构化引用。"""
    warnings = []
    table_bounds = list(_iter_table_bounds(ws))
    if not table_bounds:
        return warnings

    for (row, col), formula in formulas.items():
        if _STRUCTURED_REF_PATTERN.search(formula):
            continue
        for match in _AGG_A1_PATTERN.finditer(formula):
            sheet_name = _normalize_sheet_name(match.group("sheet"))
            if sheet_name is not None and sheet_name != ws.title:
                continue
            c1 = match.group("c1").upper()
            c2 = match.group("c2").upper()
            r1 = int(match.group("r1"))
            r2 = int(match.group("r2"))
            formula_bounds = (
                min(r1, r2),
                max(r1, r2),
                column_index_from_string(c1),
                column_index_from_string(c2),
            )
            for table_name, table_ref, t_bounds in table_bounds:
                if not _ranges_overlap(formula_bounds, t_bounds):
                    continue
                func_name = match.group(1).upper()
                warnings.append(
                    {
                        "type": "fixed_range_on_table",
                        "cell": f"{ws.title}!{get_column_letter(col)}{row}",
                        "formula": formula,
                        "table": table_name,
                        "message": (
                            f"聚合公式使用固定范围，但与 Table `{table_name}` "
                            f"（{table_ref}）重叠；应改用结构化引用如 "
                            f"`={func_name}({table_name}[列名])`（写在 Table 外）"
                        ),
                    }
                )
                break

    return warnings


def _check_structured_ref_inside_table(ws, formulas):
    """Table 内对同表列做 SUM(Table[列]) 等 → 循环引用；合计应放表外。"""
    warnings = []
    table_bounds = list(_iter_table_bounds(ws))
    if not table_bounds:
        return warnings

    for (row, col), formula in formulas.items():
        for match in _STRUCTURED_COL_AGG_PATTERN.finditer(formula):
            table_name = match.group("table")
            for name, table_ref, t_bounds in table_bounds:
                if name != table_name:
                    continue
                if not _cell_in_bounds(row, col, t_bounds):
                    continue
                func_name = match.group(1).upper()
                col_name = match.group("col")
                warnings.append(
                    {
                        "type": "structured_ref_inside_table",
                        "cell": f"{ws.title}!{get_column_letter(col)}{row}",
                        "formula": formula,
                        "table": table_name,
                        "message": (
                            f"单元格在 Table `{table_name}`（{table_ref}）内，"
                            f"却使用列聚合 `{func_name}({table_name}[{col_name}])`，"
                            "会形成循环引用；将合计写在 Table 外下一行，"
                            "或使用 Table Totals Row；行内可用 `[@列]`"
                        ),
                    }
                )
                break

    return warnings


if __name__ == "__main__":
    import sys

    if len(sys.argv) < 2:
        print("Usage: python audit.py <filename>")
        sys.exit(1)

    result = audit_formulas(sys.argv[1])
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if result.get("status") != "success":
        sys.exit(1)

