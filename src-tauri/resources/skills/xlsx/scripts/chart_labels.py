"""Chart helpers for openpyxl (Excel / WPS compatible)."""

from __future__ import annotations

import re
import warnings
import zipfile
from io import BytesIO
from pathlib import Path

from openpyxl.chart import Reference, Series
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties

# WPS needs one series per point for labels; large N bloats chart XML.
SCATTER_LABELED_POINTS_MAX = 30

_A_NS = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
_NOFILL_LN = f'<a:ln {_A_NS}><a:noFill/></a:ln>'
_SER_RE = re.compile(r'<ser>.*?</ser>', re.DOTALL)
_DLBLS_RE = re.compile(r'<dLbls>.*?</dLbls>', re.DOTALL)
_LEGEND_RE = re.compile(r'<legend>.*?</legend>', re.DOTALL)


def configure_scatter_chart(scatter) -> None:
    """Correct axPos and marker-only style (no connecting lines)."""
    scatter.x_axis.axPos = 'b'
    scatter.y_axis.axPos = 'l'
    scatter.scatterStyle = 'marker'


configure_scatter_axes = configure_scatter_chart


def _coerce_float(value) -> float | None:
    if value is None:
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _read_column_values(ws, col: int, min_row: int, max_row: int) -> list[float]:
    values: list[float] = []
    for row in range(min_row, max_row + 1):
        v = _coerce_float(ws.cell(row=row, column=col).value)
        if v is not None:
            values.append(v)
    return values


def axis_range(
    values: list[float],
    *,
    floor_zero: bool = False,
    pad_ratio: float = 0.1,
) -> tuple[float, float]:
    """Compute axis bounds with padding; optional floor at zero."""
    if not values:
        return 0.0, 1.0
    lo, hi = min(values), max(values)
    if hi == lo:
        span = abs(hi) or 1.0
        margin = span * pad_ratio
    else:
        margin = (hi - lo) * pad_ratio
    lo_out = lo - margin
    hi_out = hi + margin
    if floor_zero and min(values) >= 0:
        lo_out = 0.0
    elif min(values) >= 0:
        lo_out = max(0.0, lo_out)
    return lo_out, hi_out


def set_scatter_axis_scale(
    scatter,
    ws,
    *,
    x_col: int,
    y_col: int,
    min_row: int,
    max_row: int,
    pad_ratio: float = 0.1,
    x_floor_zero: bool = True,
    y_floor_zero: bool = False,
) -> None:
    """Set min/max and fix X-axis crossing when Y contains negatives.

    WPS/Excel default crosses X at y=0 (autoZero), which lifts the horizontal
    axis into the plot when growth rates etc. are negative. Use crosses='min'
    to keep the X-axis at the bottom of the plot area.
    """
    x_vals = _read_column_values(ws, x_col, min_row, max_row)
    y_vals = _read_column_values(ws, y_col, min_row, max_row)
    if not x_vals or not y_vals:
        return

    y_min_raw = min(y_vals)
    x_lo, x_hi = axis_range(x_vals, floor_zero=x_floor_zero, pad_ratio=pad_ratio)
    y_lo, y_hi = axis_range(
        y_vals,
        floor_zero=y_floor_zero and y_min_raw >= 0,
        pad_ratio=pad_ratio,
    )

    scatter.x_axis.scaling.min = x_lo
    scatter.x_axis.scaling.max = x_hi
    scatter.y_axis.scaling.min = y_lo
    scatter.y_axis.scaling.max = y_hi

    if y_min_raw < 0:
        scatter.x_axis.crosses = 'min'
    elif y_floor_zero:
        scatter.y_axis.scaling.min = 0


def configure_bar_chart(chart, *, horizontal: bool = True) -> None:
    """Fix axPos for bar/column charts (WPS swaps titles if left default).

    openpyxl maps x_axis → catAx (categories), y_axis → valAx (values), which is
    the opposite of intuitive "x=horizontal value, y=vertical category" naming.
    """
    if horizontal:
        chart.type = 'bar'
        chart.x_axis.axPos = 'l'  # catAx: company names on the left
        chart.y_axis.axPos = 'b'  # valAx: numeric scale on the bottom
    else:
        chart.type = 'col'
        chart.x_axis.axPos = 'b'  # catAx: categories on the bottom
        chart.y_axis.axPos = 'l'  # valAx: values on the left


def set_bar_axis_titles(chart, *, category: str, value: str) -> None:
    """Set category/value axis titles (not x/y intuitive names).

    For horizontal bars: category='公司' on the left, value='市值' on the bottom.
    """
    chart.x_axis.title = category
    chart.y_axis.title = value


def configure_scatter_series(series, *, marker_color: str | None = '4472C4', marker_size: int = 8) -> None:
    """Disable connector lines; keep visible circle markers."""
    from openpyxl.chart.marker import Marker

    no_line = GraphicalProperties()
    no_line.line = LineProperties(noFill=True)
    series.graphicalProperties = no_line

    if series.marker is None:
        series.marker = Marker(symbol='circle', size=marker_size)
    else:
        series.marker.symbol = 'circle'
        if not series.marker.size:
            series.marker.size = marker_size

    if marker_color:
        marker_gp = GraphicalProperties()
        marker_gp.solidFill = marker_color
        marker_gp.line = LineProperties(noFill=True)
        series.marker.graphicalProperties = marker_gp


def add_scatter_xy_series(
    scatter,
    ws,
    *,
    x_col: int,
    y_col: int,
    min_row: int,
    max_row: int,
    title: str | None = None,
    label_col: int | None = None,
    label_position: str = 'r',
    auto_scale: bool = True,
    x_floor_zero: bool = True,
    y_floor_zero: bool = True,
):
    """x_col → horizontal (xVal), y_col → vertical (yVal).

    Labeled scatter (label_col set): point-side labels, no legend — use scatter.title + axis titles.
    """
    if label_col is not None:
        series_list = add_scatter_labeled_points(
            scatter, ws,
            x_col=x_col, y_col=y_col, label_col=label_col,
            min_row=min_row, max_row=max_row,
            position=label_position,
        )
        scatter.legend = None
    else:
        x_ref = Reference(ws, min_col=x_col, min_row=min_row, max_row=max_row)
        y_ref = Reference(ws, min_col=y_col, min_row=min_row, max_row=max_row)
        series = Series(y_ref, x_ref, title=title)
        configure_scatter_series(series)
        scatter.series.append(series)
        series_list = series
        scatter.legend = None
        if title:
            warnings.warn(
                f'scatter title={title!r} sets legend only; '
                'point-side labels need label_col=<label column>',
                stacklevel=2,
            )

    if auto_scale:
        set_scatter_axis_scale(
            scatter, ws,
            x_col=x_col, y_col=y_col, min_row=min_row, max_row=max_row,
            x_floor_zero=x_floor_zero, y_floor_zero=y_floor_zero,
        )
    return series_list


def add_scatter_labeled_points(
    scatter,
    ws,
    *,
    x_col: int,
    y_col: int,
    label_col: int,
    min_row: int,
    max_row: int,
    position: str = 'r',
) -> list:
    """One series per point; series title becomes the visible label (WPS-safe)."""
    point_count = max_row - min_row + 1
    if point_count > SCATTER_LABELED_POINTS_MAX:
        warnings.warn(
            f'scatter labeled points={point_count} exceeds recommended max '
            f'{SCATTER_LABELED_POINTS_MAX}; chart XML may be large and slow in WPS. '
            'Consider fewer points or no per-point labels.',
            stacklevel=2,
        )

    series_list = []
    for row in range(min_row, max_row + 1):
        label = ws.cell(row=row, column=label_col).value
        if label is None:
            label = ''
        x_ref = Reference(ws, min_col=x_col, min_row=row, max_row=row)
        y_ref = Reference(ws, min_col=y_col, min_row=row, max_row=row)
        series = Series(y_ref, x_ref, title=str(label))
        configure_scatter_series(series)
        series.dLbls = DataLabelList()
        series.dLbls.showSerName = True
        series.dLbls.showVal = False
        series.dLbls.showCatName = False
        series.dLbls.showLegendKey = False  # WPS renders tiny legend-key circle otherwise
        series.dLbls.showLeaderLines = False
        series.dLbls.dLblPos = position
        scatter.series.append(series)
        series_list.append(series)

    return series_list


def finalize_scatter_chart(xlsx_path: str) -> None:
    """Post-save patch: marker-only, no connector lines, hide multi-series legend (WPS)."""
    patch_scatter_no_lines(xlsx_path)


def _patch_single_ser(ser_xml: str) -> str:
    """Patch one <ser> block: connector noFill, marker circle, dLbls flags."""
    if '<spPr>' in ser_xml:
        ser_xml = re.sub(
            r'(<spPr>)(<a:ln[^>]*>.*?</a:ln>)',
            rf'\1{_NOFILL_LN}',
            ser_xml,
            count=1,
            flags=re.DOTALL,
        )

    def _fix_marker(m: re.Match[str]) -> str:
        block = m.group(0)
        block = re.sub(r'(<symbol val=")none(")', r'\1circle\2', block)
        inner = block[len('<marker>'):block.rfind('</marker>')]
        if 'symbol val=' not in inner:
            block = block.replace('<marker>', '<marker><symbol val="circle"/><size val="8"/>', 1)
        else:
            block = re.sub(
                r'(<symbol val="circle"/><size val=")\d+(")',
                r'\g<1>8\2',
                block,
                count=1,
            )
        return block

    ser_xml = re.sub(r'<marker>.*?</marker>', _fix_marker, ser_xml, flags=re.DOTALL)

    def _fix_dlbls(m: re.Match[str]) -> str:
        block = m.group(0)
        if 'showLegendKey' in block:
            block = re.sub(r'<showLegendKey val="[^"]*"/>', '<showLegendKey val="0"/>', block)
        else:
            block = block.replace('<dLbls>', '<dLbls><showLegendKey val="0"/>', 1)
        return re.sub(r'<showLeaderLines val="1"/>', '<showLeaderLines val="0"/>', block)

    ser_xml = _DLBLS_RE.sub(_fix_dlbls, ser_xml)
    return ser_xml


def _plot_area_layout_xml() -> str:
    return (
        '<layout><manualLayout>'
        '<layoutTarget val="inner"/>'
        '<xMode val="edge"/><yMode val="edge"/>'
        '<x val="0.12"/><y val="0.14"/>'
        '<w val="0.83"/><h val="0.64"/>'
        '</manualLayout></layout>'
    )


def _patch_chart_spacing(xml: str) -> str:
    """Reserve margins so chart title and axis titles do not overlap tick labels."""
    plot_layout = _plot_area_layout_xml()
    if re.search(r'<plotArea>\s*<layout>', xml):
        xml = re.sub(
            r'(<plotArea>)\s*<layout>.*?</layout>',
            rf'\1{plot_layout}',
            xml,
            count=1,
            flags=re.DOTALL,
        )
    else:
        xml = xml.replace('<plotArea>', f'<plotArea>{plot_layout}', 1)

    plot_idx = xml.find('<plotArea>')
    if plot_idx > 0:
        head, tail = xml[:plot_idx], xml[plot_idx:]

        def _title_layout(m: re.Match[str]) -> str:
            block = m.group(0)
            if '<layout>' in block:
                return block
            extra = (
                '<layout><manualLayout>'
                '<xMode val="edge"/><yMode val="edge"/>'
                '<x val="0"/><y val="0"/><w val="1"/><h val="0.12"/>'
                '</manualLayout></layout><overlay val="0"/>'
            )
            return block.replace('</tx>', f'</tx>{extra}', 1)

        head = re.sub(r'<title>.*?</title>', _title_layout, head, count=1, flags=re.DOTALL)
        xml = head + tail

    def _patch_bottom_axis(m: re.Match[str]) -> str:
        block = m.group(0)
        if 'lblOffset' not in block:
            block = block.replace('<majorGridlines/>', '<majorGridlines/><lblOffset val="150"/>', 1)
        if 'tickLblPos' not in block:
            block = block.replace('<crossAx', '<tickLblPos val="low"/><crossAx', 1)
        if '<title>' in block and '<overlay' not in block:
            block = re.sub(r'(</title>)', r'<overlay val="0"/>\1', block, count=1)
        return block

    xml = re.sub(
        r'<valAx>(?:(?!</valAx>).)*<axPos val="b"/>(?:(?!</valAx>).)*</valAx>',
        _patch_bottom_axis,
        xml,
        flags=re.DOTALL,
    )
    return xml


def _suppress_multi_series_legend(xml: str, ser_count: int) -> str:
    """Hide sidebar legend for labeled scatter (WPS lists every series if legend is empty)."""
    xml = _LEGEND_RE.sub('', xml)
    if ser_count <= 1:
        return xml

    # WPS ignores off-screen manualLayout; must mark every entry deleted.
    entries = ''.join(
        f'<legendEntry><idx val="{i}"/><delete val="1"/></legendEntry>'
        for i in range(ser_count)
    )
    legend_xml = (
        f'<legend><legendPos val="r"/><overlay val="1"/>'
        f'{entries}'
        f'<layout><manualLayout>'
        f'<layoutTarget val="inner"/>'
        f'<xMode val="edge"/><yMode val="edge"/>'
        f'<x val="1"/><y val="0"/><w val="0"/><h val="0"/>'
        f'</manualLayout></layout>'
        f'</legend>'
    )
    xml = xml.replace('</plotArea>', f'</plotArea>{legend_xml}', 1)
    return xml


def _patch_scatter_chart_xml(xml: str) -> tuple[str, int]:
    """Return patched XML and number of <ser> blocks processed."""
    if 'scatterStyle' in xml:
        xml = re.sub(r'<scatterStyle val="[^"]*"/>', '<scatterStyle val="marker"/>', xml)
    else:
        xml = xml.replace('<scatterChart>', '<scatterChart><scatterStyle val="marker"/>', 1)

    ser_blocks = _SER_RE.findall(xml)
    ser_count = len(ser_blocks)
    patched = [_patch_single_ser(ser) for ser in ser_blocks]
    it = iter(patched)
    xml = _SER_RE.sub(lambda _m: next(it), xml)
    xml = re.sub(r'<smooth val="1"/>', '', xml)

    if 'scatterChart' in xml and 'scatterStyle val="marker"' not in xml:
        warnings.warn('patch_scatter_no_lines: scatterStyle marker not applied', stacklevel=2)
    if ser_count == 0 and 'scatterChart' in xml:
        warnings.warn('patch_scatter_no_lines: scatterChart found but no <ser> patched', stacklevel=2)

    xml = _patch_chart_spacing(xml)
    xml = _suppress_multi_series_legend(xml, ser_count)
    return xml, ser_count


def patch_scatter_no_lines(xlsx_path: str) -> None:
    """Force scatterStyle=marker and noFill on series connector lines."""
    path = Path(xlsx_path)
    with zipfile.ZipFile(path, 'r') as zin:
        names = zin.namelist()
        contents = {n: zin.read(n) for n in names}

    changed = False
    for chart_path in names:
        if not re.match(r'xl/charts/chart\d+\.xml$', chart_path):
            continue
        xml = contents[chart_path].decode('utf-8')
        if 'scatterChart' not in xml:
            continue

        xml, _ = _patch_scatter_chart_xml(xml)
        contents[chart_path] = xml.encode('utf-8')
        changed = True

    if not changed:
        return

    out = BytesIO()
    with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as zout:
        for name in names:
            zout.writestr(name, contents[name])
    path.write_bytes(out.getvalue())
